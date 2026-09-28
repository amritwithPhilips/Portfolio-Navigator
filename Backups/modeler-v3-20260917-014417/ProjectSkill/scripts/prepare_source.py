"""Normalize the authoritative AOP 2027 workbook for the model and dashboard."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").replace("\u200b", " ").split())


def number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    return None


def normalize(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []

    for sheet_name in ("EW", "Bu"):
        sheet = workbook[sheet_name]
        headers = [clean(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        header_index = {header: index for index, header in enumerate(headers) if header}
        decision_band_index = 50 if sheet_name == "EW" else 48
        overall_decision_index = 64 if sheet_name == "EW" else 62
        priority_index = 21 if sheet_name == "EW" else 19

        for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            item = clean(row[0])
            name = clean(row[3])
            requested_k = number(row[14])
            if not any((item, name, requested_k is not None)):
                continue
            record_id = item or f"{sheet_name}-row-{row_number}"
            benefit_k = number(row[16])
            record = {
                "id": record_id,
                "item": item,
                "temporaryId": not bool(item),
                "planisware": clean(row[1]),
                "sourceSheet": sheet_name,
                "sourceRow": row_number,
                "stream": "ew" if sheet_name == "EW" else "bu",
                "org": clean(row[2]),
                "name": name,
                "newCarry": clean(row[5]),
                "category": clean(row[6]),
                "platform": clean(row[7]),
                "costM": round((requested_k or 0) / 1000, 6),
                "requestedRaw": row[14],
                "benefitM": round(benefit_k / 1000, 6) if benefit_k is not None else None,
                "benefitRaw": row[16],
                "priority": clean(row[priority_index]),
                "impactRaw": clean(row[header_index["Impact and Cost of Not Doing"]]),
                "problem": clean(row[header_index["Problem Statement"]]),
                "scope": clean(row[header_index["Scope"]]),
                "dependency": clean(row[header_index["Dependency (Item No & Explaination)"]] if sheet_name == "EW" else row[header_index["Dependency"]]),
                "decisionBand": clean(row[decision_band_index]),
                "overallDecision": clean(row[overall_decision_index]),
            }
            records.append(record)

    ids = [record["id"] for record in records]
    duplicates = sorted({record_id for record_id in ids if ids.count(record_id) > 1})
    if duplicates:
        raise ValueError(f"Duplicate normalized IDs: {duplicates}")

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    totals = {
        "records": len(records),
        "itemReferences": sum(bool(record["item"]) for record in records),
        "requestedM": round(sum(record["costM"] for record in records), 6),
        "ewRequestedM": round(sum(record["costM"] for record in records if record["stream"] == "ew"), 6),
        "buRequestedM": round(sum(record["costM"] for record in records if record["stream"] == "bu"), 6),
    }
    return {
        "sourceFile": source.name,
        "sourceSha256": digest,
        "totals": totals,
        "records": records,
    }


def write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(records[0]) if records else []
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--json", type=Path)
    parser.add_argument("--csv", type=Path)
    args = parser.parse_args()

    payload = normalize(args.source)
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.csv:
        write_csv(args.csv, payload["records"])
    print(json.dumps(payload["totals"], ensure_ascii=False))


if __name__ == "__main__":
    main()
