import json
from pathlib import Path
from openpyxl import load_workbook
D=Path(__file__).parent/'audit';checks=[]
def ck(name,passed,detail=''):checks.append(dict(name=name,passed=bool(passed),detail=detail))
def near(a,b):return abs((a or 0)-(b or 0))<1e-8
for case in json.loads((D/'rebuild-downloads.json').read_text()):
 w=load_workbook(D/(case['id']+'.xlsx'),read_only=True,data_only=True)
 meta={r[0]:r[1] for r in w['Scenario Selection'].values}
 rows=list(w['Selected Programs'].values);records=[dict(zip(rows[0],r)) for r in rows[1:]];by={r['ID']:r for r in records}
 ck(case['id']+' exact scenario revision',meta['Scenario ID']==case['scenarioId'] and meta['Revision']==case['revision'] and meta['Scenario']==case['name'])
 ck(case['id']+' scope and row IDs',meta['Export scope']==case['scope'] and set(by)=={r['id'] for r in case['rows']} and len(by)==len(records))
 ck(case['id']+' exact request funding merit',all(near(by[r['id']]['Requested M€'],r['cost']) and near(by[r['id']]['Funded M€'],r['final']) and near(by[r['id']]['Merit'],r['score']) for r in case['rows']))
 ck(case['id']+' company finance and selected totals',near(meta['Company funded base M€'],case['base']) and near(meta['Company total IT M€'],case['total']) and near(meta['Selected funded M€'],sum(r['final'] for r in case['rows'])))
 ck(case['id']+' original source references',all(r['Source sheet'] in ['EW','BU'] and isinstance(r['Source row'],int) for r in records))
 ck(case['id']+' unambiguous filename','-r'+str(case['revision'])+'-'+case['scope']+'.xlsx' in case['filename'])
 ck(case['id']+' finance and decision evidence',all(n in w.sheetnames for n in ['2. Programs','9. Source Reference']) and any('Ledger' in n for n in w.sheetnames))
 w.close()
(D/'scenario-download-qa.json').write_text(json.dumps({'checks':checks},indent=2))
print(json.dumps({'checks':len(checks),'passed':sum(c['passed'] for c in checks),'failures':[c for c in checks if not c['passed']]}))
