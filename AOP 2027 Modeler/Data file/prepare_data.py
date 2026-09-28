"""
Prepare demand data JSON from any AOP 2027 demand XLSX.

Reads the XLSX via openpyxl, validates structure, produces a workbook-format
JSON that the browser engine can normalize directly. Runs QA checks and
embeds the report in the output.

Usage:
  python prepare_data.py
  python prepare_data.py --source "path/to/demand.xlsx"
  python prepare_data.py --source "path/to/demand.xlsx" --output "AOP2027-Demand-Data.json"
"""
import argparse, base64, datetime, hashlib, io, json, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
from openpyxl import load_workbook

DEFAULT_SOURCE = Path(r'C:\Users\320126040\Documents\Codex\AOP2027\AOP 2027 Demand\AOP 2027 Consolidated Live.xlsx')
DEFAULT_OUTPUT = Path(__file__).resolve().parent / 'AOP2027-Demand-Data.json'

REQUIRED_SHEETS = ['EW', 'BU']
REQUIRED_COLUMNS = [
    'Item No.', 'Planisware ID', 'Initiative name', 'Type',
    'Category', 'Leading IT Platform', 'Investment Area', 'Priority',
    'Overall AOP decision', 'Scope', 'Problem Statement',
    'Impact and Cost of Not Doing',
]
PARTIAL_MATCH_COLUMNS = [
    'Organization', 'Carry-over', 'IT Investment', 'Quantified Benefits', 'Payback',
]

BASELINE = {'rows': 509, 'demand_m': 279.595, 'ew_rows': 180, 'bu_rows': 329}
DRIFT_TOLERANCE_ROWS = 0.5   # allow up to 50% drift from baseline row count
DRIFT_TOLERANCE_DEMAND = 0.5  # allow up to 50% drift from baseline demand


def safe(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def find_col(headers, pattern):
    pattern_l = pattern.lower()
    for i, h in enumerate(headers):
        if h and pattern_l in ' '.join(str(h).split()).lower():
            return i
    return -1


def find_header_row(rows, marker='Item No.', scan_limit=10):
    """Auto-detect the header row by scanning the first *scan_limit* rows
    for one that contains *marker* (case-insensitive).  Returns the 0-based
    index, or 0 as a fallback."""
    marker_l = marker.lower()
    for idx in range(min(scan_limit, len(rows))):
        for cell in rows[idx]:
            if cell and marker_l in str(cell).lower():
                return idx
    return 0


def run_qa(sheets_data, raw_bytes):
    checks = []
    summary = {
        'rows': 0, 'ewRows': 0, 'buRows': 0,
        'totalDemandK': 0, 'hpmRows': 0, 'hpmDemandK': 0,
        'bandsFilled': 0, 'bandsBlank': 0,
        'negativeCosts': 0, 'duplicateIds': [],
    }

    def add(name, severity, status, detail):
        checks.append({'name': name, 'severity': severity, 'status': status, 'detail': detail})

    sheet_names = [s['name'] for s in sheets_data]
    for required in REQUIRED_SHEETS:
        found = any(s.lower() == required.lower() for s in sheet_names)
        add(f'Sheet {required} exists', 'critical', 'PASS' if found else 'FAIL',
            f'Found' if found else f'Sheet named "{required}" not found')

    if not all(any(s.lower() == r.lower() for s in sheet_names) for r in REQUIRED_SHEETS):
        has_fail = any(c['status'] == 'FAIL' for c in checks)
        return {'status': 'FAIL' if has_fail else 'PASS', 'checks': checks, 'summary': summary}

    ids_seen = set()
    dupes = []
    missing_cols = []

    required_lower = [r.lower() for r in REQUIRED_SHEETS]
    for sheet in sheets_data:
        sname = sheet['name']
        if sname.lower() not in required_lower:
            continue
        rows = sheet['rows']
        if len(rows) < 2:
            add(f'{sname}: has data rows', 'critical', 'FAIL', f'Only {len(rows)} rows (need header + data)')
            continue

        header_row = find_header_row(rows)
        headers = [str(h) if h else '' for h in rows[header_row]]

        for col_name in REQUIRED_COLUMNS:
            if find_col(headers, col_name) == -1 and col_name not in missing_cols:
                missing_cols.append(col_name)
        for pattern in PARTIAL_MATCH_COLUMNS:
            if find_col(headers, pattern) == -1 and pattern not in missing_cols:
                missing_cols.append(pattern)

        sub_row = header_row + 1
        sub_headers = [str(h) if h else '' for h in rows[sub_row]] if sub_row < len(rows) else []
        band_idx = find_col(headers, 'Decision Band')
        if band_idx == -1:
            for i, (h, sh) in enumerate(zip(headers, sub_headers + [''] * len(headers))):
                hl, sl = h.lower(), sh.lower()
                if 'waste' in hl or 'decision band' in hl or 'band' in sl:
                    band_idx = i
                    break

        cost_idx = find_col(headers, 'IT Investment')
        item_idx = find_col(headers, 'Item No.')
        org_idx = find_col(headers, 'Organization')
        name_idx = find_col(headers, 'Initiative name')

        # Start data rows after the header (skip any sub-header rows too)
        data_start = header_row + 1
        # If there's a sub-header row that doesn't look like data, skip it
        if data_start < len(rows) and sub_headers:
            first_item = ''
            if item_idx >= 0 and item_idx < len(sub_headers):
                first_item = sub_headers[item_idx].strip()
            # Sub-headers typically don't have numeric Item No. values
            if first_item and not any(c.isdigit() for c in first_item):
                data_start = header_row + 2

        for i in range(data_start, len(rows)):
            row = rows[i]
            item = str(row[item_idx]).strip() if item_idx >= 0 and item_idx < len(row) and row[item_idx] else ''
            init_name = str(row[name_idx]).strip() if name_idx >= 0 and name_idx < len(row) and row[name_idx] else ''
            cost_raw = row[cost_idx] if cost_idx >= 0 and cost_idx < len(row) else None

            cost_val = None
            if cost_raw is not None and cost_raw != '' and cost_raw is not None:
                try:
                    cost_val = float(str(cost_raw).replace('€', '').replace(',', '').strip())
                except (ValueError, TypeError):
                    pass

            if not item and not init_name and cost_val is None:
                continue

            summary['rows'] += 1
            stream = 'ew' if sname.lower() == 'ew' else 'bu'
            if stream == 'ew':
                summary['ewRows'] += 1
            else:
                summary['buRows'] += 1

            row_id = item or f'{sname}-row-{i+1}'
            if row_id in ids_seen:
                dupes.append(row_id)
            ids_seen.add(row_id)

            if cost_val is not None:
                if cost_val < 0:
                    summary['negativeCosts'] += 1
                summary['totalDemandK'] += cost_val

            org = ''
            if org_idx >= 0 and org_idx < len(row) and row[org_idx]:
                org = str(row[org_idx]).strip()
            if 'hpm' in org.lower():
                summary['hpmRows'] += 1
                summary['hpmDemandK'] += (cost_val or 0)

            if band_idx >= 0 and band_idx < len(row) and row[band_idx] and str(row[band_idx]).strip():
                summary['bandsFilled'] += 1
            else:
                summary['bandsBlank'] += 1

    summary['duplicateIds'] = dupes
    total_demand_m = round(summary['totalDemandK'] / 1000, 3)
    hpm_demand_m = round(summary['hpmDemandK'] / 1000, 3)

    add('Required columns present', 'critical', 'PASS' if not missing_cols else 'FAIL',
        f'All matched' if not missing_cols else f'Missing: {", ".join(missing_cols)}')
    add('Data rows > 0', 'critical', 'PASS' if summary['rows'] > 0 else 'FAIL',
        f'{summary["rows"]} total data rows')

    row_drift = abs(summary['rows'] - BASELINE['rows']) / BASELINE['rows'] if BASELINE['rows'] > 0 else 0
    add('Row count within range', 'warning', 'PASS' if row_drift <= DRIFT_TOLERANCE_ROWS else 'WARN',
        f'{summary["rows"]} rows (baseline {BASELINE["rows"]}, {row_drift*100:.1f}% drift)')

    demand_drift = abs(total_demand_m - BASELINE['demand_m']) / BASELINE['demand_m'] if BASELINE['demand_m'] > 0 else 0
    add('Total demand within range', 'warning', 'PASS' if demand_drift <= DRIFT_TOLERANCE_DEMAND else 'WARN',
        f'{total_demand_m:.3f} M€ (baseline {BASELINE["demand_m"]:.3f} M€, {demand_drift*100:.1f}% drift)')

    add('HPM rows present', 'warning', 'PASS' if summary['hpmRows'] > 0 else 'WARN',
        f'{summary["hpmRows"]} HPM rows, {hpm_demand_m:.3f} M€')

    band_pct = summary['bandsFilled'] / summary['rows'] if summary['rows'] > 0 else 0
    add('Decision bands populated', 'warning', 'PASS' if band_pct >= 0.5 else 'WARN',
        f'{summary["bandsFilled"]} of {summary["rows"]} rows ({band_pct*100:.0f}%)')

    add('No negative costs', 'critical', 'PASS' if summary['negativeCosts'] == 0 else 'FAIL',
        'All costs >= 0' if summary['negativeCosts'] == 0 else f'{summary["negativeCosts"]} rows have negative costs')

    add('No duplicate Item No.', 'critical', 'PASS' if not dupes else 'FAIL',
        f'{len(ids_seen)} unique IDs' if not dupes else f'Duplicates: {", ".join(dupes[:10])}')

    lfr_ids = ['EW-11', 'EW-65', 'EW-71']
    found_lfr = [x for x in lfr_ids if x in ids_seen]
    add('LFR exceptions present', 'info', 'PASS' if len(found_lfr) == 3 else 'WARN',
        f'{len(found_lfr)} of 3 ({", ".join(lfr_ids)})')

    has_fail = any(c['status'] == 'FAIL' for c in checks)
    has_warn = any(c['status'] == 'WARN' for c in checks)

    return {
        'status': 'FAIL' if has_fail else ('WARN' if has_warn else 'PASS'),
        'checks': checks,
        'summary': {
            'rows': summary['rows'],
            'ewRows': summary['ewRows'],
            'buRows': summary['buRows'],
            'totalDemand': total_demand_m,
            'hpmRows': summary['hpmRows'],
            'hpmDemand': hpm_demand_m,
            'bandsFilled': summary['bandsFilled'],
            'bandsBlank': summary['bandsBlank'],
        }
    }


def extract_sheets(source_path):
    wb = load_workbook(source_path, read_only=True, data_only=True)
    sheets = []
    for ws in wb:
        rows = [[safe(cell) for cell in row] for row in ws.iter_rows(values_only=True)]
        sheets.append({'name': ws.title, 'rows': rows})
    wb.close()
    return sheets


def main():
    parser = argparse.ArgumentParser(description='Prepare AOP 2027 demand data JSON')
    parser.add_argument('--source', type=Path, default=DEFAULT_SOURCE, help='Path to demand XLSX')
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT, help='Output JSON path')
    args = parser.parse_args()

    source = args.source
    if not source.exists():
        print(f'ERROR: Source file not found: {source}', file=sys.stderr)
        sys.exit(1)

    print(f'Reading {source.name}...')
    raw = source.read_bytes()
    source_hash = hashlib.sha256(raw).hexdigest()
    original_b64 = base64.b64encode(raw).decode()

    print(f'  Size: {len(raw):,} bytes')
    print(f'  SHA-256: {source_hash[:24]}...')

    sheets = extract_sheets(source)
    print(f'  Sheets: {", ".join(s["name"]+" ("+str(len(s["rows"]))+" rows)" for s in sheets)}')

    qa = run_qa(sheets, raw)
    s = qa['summary']

    print(f'\nQA Status: {qa["status"]}')
    for check in qa['checks']:
        icon = '✓' if check['status'] == 'PASS' else ('⚠' if check['status'] == 'WARN' else '✗')
        print(f'  {icon} [{check["severity"]:8}] {check["name"]}: {check["detail"]}')

    print(f'\nSummary:')
    print(f'  Total rows: {s["rows"]} (EW {s["ewRows"]}, BU {s["buRows"]})')
    print(f'  Total demand: {s["totalDemand"]:.3f} M€')
    print(f'  HPM: {s["hpmRows"]} rows, {s["hpmDemand"]:.3f} M€')
    print(f'  Decision bands: {s["bandsFilled"]} filled, {s["bandsBlank"]} blank')

    output = {
        'schema': 4,
        'sourceFile': source.name,
        'sourceHash': source_hash,
        'preparedAt': datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'qa': qa,
        'sheets': [{'name': s['name'], 'rows': s['rows']} for s in sheets if s['name'] in REQUIRED_SHEETS or s['name'].lower() in [x.lower() for x in REQUIRED_SHEETS]],
        'originalBase64': original_b64,
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with open(args.output, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, separators=(',', ':'))

    size_mb = args.output.stat().st_size / (1024 * 1024)
    print(f'\nWritten: {args.output} ({size_mb:.1f} MB)')
    print(f'Pair with the dashboard HTML to serve.')

    if qa['status'] == 'FAIL':
        print(f'\nWARNING: QA FAILED — fix critical issues before using this data.', file=sys.stderr)
        sys.exit(2)


if __name__ == '__main__':
    main()
