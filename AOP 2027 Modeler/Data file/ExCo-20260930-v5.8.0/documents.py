from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
D=Path(__file__).parent;req=json.loads((D/'requirements.json').read_text(encoding='utf8'))
def new(title):
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8)
 for n in ['Normal','Title','Heading 1','Heading 2','Header','Footer']:
  st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(8);d.styles['Normal'].paragraph_format.line_spacing=1.12;d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
 s.header.paragraphs[0].text='AOP 2027  |  ExCo decision view  |  Version 5.8.0'
 f=s.footer.paragraphs[0];f.text='30 September 2026  |  ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 for st in d.styles:
  for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 d.add_paragraph(title,'Title');return d
def para(d,text):d.add_paragraph(text)
def h(d,text):d.add_paragraph(text,'Heading 1')
d=new('AOP 2027 ExCo view product requirements')
para(d,'Version 5.8.0 addendum. This specification adds the leadership decision view to the six-tab portfolio modeler. It supplements AOP2027-Modeler-PRD-v5.7.0.docx; the earlier document remains the specification for unchanged detailed workflows. The v5.7.0 HTML and both source workbooks are preserved.')
h(d,'Product outcome')
para(d,'ExCo can assess company IT cost, inspect accountable Programs, compare a recommended scenario and preview funding decisions in one concise view. Portfolio analysts can open the existing six tabs to change advanced configurations and review source evidence. Both views share one scenario and finance ledger.')
h(d,'Meeting journey')
for title,text in [('Company position','Five KPIs show total IT, base Programs, budget headroom or excess, funded Programs and unallocated cost. Priority funding and receiving organisation costs explain the company position.'),('My portfolio','Select one of 16 ExCo leaders to inspect accountable Programs. Keep hard and soft submitted lifetime benefit claims separate. Open a Programs ID for source evidence.'),('Compare options','Compare baseline, recommended and current funding. Show the Programs that change from the current scenario to the recommendation. Save the current scenario before creating a recommended child draft.'),('Decisions','Preview a selected leader and stream trade-off, or a discretionary Programs scope decision with rationale. Review exact cost changes before applying. Save and export the meeting outcome.')]:
 d.add_paragraph(title,'Heading 2');para(d,text)
d.add_page_break();h(d,'Recommendation and finance contract')
para(d,'Retain the default €100M Optimal scenario as the decision baseline. It is the best of 135 tested lever configurations under an objective that protects baseline-funded LFR and confirmed carry-over, retains at least €20.609M Sustain and maximises funding-weighted source merit. This is a finite scenario search, not proof of a global company optimum.')
para(d,'Recommended base Programs funding is €99.963M: LFR €43.776M, confirmed carry-over €16.773M, Sustain €20.609M and Value €18.805M. PMI adds €5.19M and Operations adds €470.314463039M, producing total IT cost of €575.467463039M. Default tolerance is zero; source merit floor 45 and weights 40/25/15/10/10 are retained.')
para(d,'BU row 276 Enterprise-wide charging is a percentage key only; its €40M is excluded from demand. Apply the approved 1,000× fraction correction and retain 22.6303317536% of EW unallocated. Optimal has €11.250895735M EW unallocated and €1.822M other BU unallocated. PMI has no BU charge. Operations own, Region and Function contributions are counted once.')
para(d,'Confirmed carry-over funding uses the lesser of the requested amount and approved baseline cap. No carry-over ROI threshold is introduced. Validated benefit evidence remains eligible; lifetime claims are not presented as annual savings. Nine missing-request and four uninterpretable-decision rows remain audit-only.')
h(d,'Decision integrity')
para(d,'Preview computes a proposed copy. Apply requires unchanged scenario ID, revision and source hash. Finalised scenarios block edits; a working version retains lineage. Manual decisions can exceed the budget and the excess remains explicit. Operations costs remain fixed to the source allocation.')
para(d,'The recommendation is bound to the Programs source hash. A different source requires a new recommendation analysis. The offline Finalised status is a local lock, not a central approval signature. Export the full evidence workbook and save the scenario for recovery.')
d.add_page_break();h(d,'Release acceptance and limitations')
para(d,'The release must pass 32 scenarios across six tabs, ExCo finance parity, actual configuration controls, 16 leader filters, decision preview and cancellation, undo, save and reload, finalise and fork, portable recovery, source refresh holds and source-linked Excel exports. An independent source replay validates 69,280 record fields, receiving organisation totals and Operations components.')
para(d,'Funding, scoring and finance results must match v5.7.0 for the same valid inputs. The only engine validation enhancement rejects invalid imported request and merit reviews. No approved finance or prioritisation rule is changed. The final validation report records actual assertion totals and evidence.')
h(d,'Implementation architecture')
para(d,'The ExCo extension reads the existing scenario result through AOPApp. Commands delegate to the existing transaction, trade-off, undo, save and export functions. A state-change event refreshes the ExCo presentation. The detailed six-tab workbench remains available without a separate calculation model.')
h(d,'Known decision gaps')
para(d,'The source does not provide a quantified dependency and delivery-capacity optimisation model. Submitted benefits can overlap and do not establish realised cash flows. €180.0599M of base request remains unfunded under the recommended rules. Receiving allocation retains €0.465M for Bu-06 without a percentage and €1.357M explicit indirect unallocated cost. These balances remain visible until a source decision changes them.')
h(d,'Companion evidence')
for t in ['ExCo-QA-report-v5.8.0.html contains all 32 scenarios, exact costs, findings and the final test score.','ExCo-architecture-v5.8.0.html contains the implemented view architecture and acceptance requirements.','32-scenarios-Programs-trace.csv contains 13,856 source-linked scenario records.','135-options-ranking.csv contains the tested recommendation candidates and settings.','AOP2027-Modeler-Detailed-Requirements-v5.8.0-ExCo-Addendum.docx lists the new acceptance contract.']:para(d,t)
d.save(D/'AOP2027-Modeler-PRD-v5.8.0-ExCo-Addendum.docx')
d=new('AOP 2027 ExCo view detailed requirements')
d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
d.styles['Heading 1'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
para(d,'Version 5.8.0 addendum. These requirements supplement AOP2027-Modeler-Detailed-Requirements-v5.7.0.docx. They cover the new ExCo presentation and command bridge. Existing funding, scoring, source exclusions and recipient allocation rules remain in force.')
for i,(ident,title,text) in enumerate(req):
 # Let requirements flow with headings kept with their acceptance text.
 d.add_paragraph(ident+'  '+title,'Heading 1');p=d.add_paragraph(text);p.paragraph_format.keep_with_next=True
 evidence={0:'ExCo default view and six-tab navigation browser checks.',1:'32-scenario ExCo parity, independent source replay and downloaded Excel ledger reconciliation.',2:'All 16 leader ownership filters plus company recipient totals.',3:'Hard and soft benefit rendering checks and source classification regression.',4:'135-option analysis and changed-source recommendation guard.',5:'Recommended child draft preserves parent ID and expected funding.',6:'Rationale, preview non-mutation, cancellation and partial-scope commit browser checks.',7:'Stale preview rejected after a concurrent revision change.',8:'F&R and BU Led trade-offs match scenario 09; base budget stays €100M.',9:'Browser save/reload, undo, finalise/fork, exports, and refreshed decision holds.',10:'Negative, non-numeric and out-of-range import rejection with valid zero and 100 boundaries.',11:'SHA256 preservation, 32-case v5.7 engine parity and full validation report.'}[i]
 para(d,'Acceptance evidence: '+evidence)
d.save(D/'AOP2027-Modeler-Detailed-Requirements-v5.8.0-ExCo-Addendum.docx')
print('Created two v5.8.0 specification addenda')
