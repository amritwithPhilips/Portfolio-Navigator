from pathlib import Path
import json, zipfile, math
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image, ImageOps, ImageDraw

root=Path(__file__).resolve().parents[2]
out=root/'output'/'modeler-v3'
expected=json.loads((out/'Data file/AOP2027-Default-Results.json').read_text(encoding='utf-8'))
wb=load_workbook(out/'Data file/2027 IT Portfolio Prioritisation Model.xlsx',data_only=True,read_only=True)
rows=list(wb['Portfolio'].values); header=rows[0]; actual={r[0]:dict(zip(header,r)) for r in rows[1:]}
assert len(actual)==509
for r in expected['rows']:
    a=actual[r['id']]
    for key,col in [('cost','Requested M€'),('automatic','Automatic M€'),('final','Prioritised M€'),('unfunded','Unfunded M€'),('score','Merit')]:
        assert a[col]==r[key] or (a[col] is not None and r[key] is not None and abs(a[col]-r[key])<1e-8),(r['id'],key,a[col],r[key])
assert sum(x['Prioritised M€'] or 0 for x in actual.values())-expected['totals']['allocation']<1e-8
for row in wb['Reconciliation'].iter_rows(min_row=2,max_row=5,min_col=4,max_col=4,values_only=True):assert abs(row[0])<1e-8,row
portable=json.loads((out/'QA-Portable.json').read_text(encoding='utf-8'));session=json.loads(portable['payload'])
browser=load_workbook(out/'QA-Export.xlsx',data_only=True,read_only=True)
exported=list(browser['Portfolio'].values);assert len(exported)==510
assert abs(sum(r[13] or 0 for r in exported[1:])-sum(v['amount'] if v['choice']=='include' else 0 for v in []))>=0 # exact validation below
base={r['id']:r for r in expected['rows']}
for row in exported[1:]:
    ov=session['overrides'].get(row[0]);wanted=ov['amount'] if ov else base[row[0]]['final']
    assert abs(row[13]-wanted)<1e-8,(row[0],row[13],wanted)
assert 'Source Reference' in browser.sheetnames
checks={'static_workbook_rows':509,'static_row_financial_parity':'PASS','reconciliation_formulas':'PASS','browser_export_rows':509,'browser_final_scenario_parity':'PASS'}
qa=root/'tmp/modeler-requirements-qa'
for stem in ['AOP2027-Modeler-PRD-v3.2','AOP2027-Modeler-Detailed-Requirements-v3.2']:
    folder=qa/stem;pdf=PdfReader(folder/(stem+'.pdf'))
    text='\n'.join(p.extract_text() for p in pdf.pages)
    assert 'Function' in text and 'BU Led' in text
    assert '46' in text and 'Enterprise Wide' in text and 'BU Led' in text
    assert 'Demand Prioritized' in text and 'Demand Not Funded' in text
    assert 'Funding Prioritization Gate' in text
    assert 'complete the conference' not in text.lower()
    assert 'review organisations in this conference' not in text.lower()
    assert not any(s in text for s in ['approval is pending','Approval status: pending','Carryover 102.0'])
    checks[stem]={'pages':len(pdf.pages),'content_checks':'PASS'}
    imgs=sorted(folder.glob('page-*.png'))
    for batch in range(0,len(imgs),6):
        thumb_w,thumb_h=435,615;canvas=Image.new('RGB',(thumb_w*3, (thumb_h+28)*2),'#dbe2ea');draw=ImageDraw.Draw(canvas)
        for j,p in enumerate(imgs[batch:batch+6]):
            im=Image.open(p).convert('RGB');im.thumbnail((thumb_w,thumb_h));x=(j%3)*thumb_w;y=(j//3)*(thumb_h+28);canvas.paste(im,(x,y+28));draw.text((x+10,y+8),f'{stem} page {batch+j+1}',fill='black')
        canvas.save(out/f'QA-{stem}-{batch//6+1}.png')
for name in ['Engine-Test-Results.json','Browser-Test-Results.json','Interaction-Test-Results-v3.2.json','Refresh-Test-Results-v3.2.json','Scenario-Outcomes-v3.2.json','Skill-Test-Results.json']:
    report=json.loads((out/name).read_text(encoding='utf-8'))
    assert report.get('failures',0)==0,(name,report.get('failures'))
    checks[name]={'failures':0,'status':'PASS'}
(out/'Package-Test-Results.json').write_text(json.dumps(checks,indent=2))
print(json.dumps(checks,indent=2))
