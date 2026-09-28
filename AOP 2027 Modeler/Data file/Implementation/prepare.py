"""Preserve the complete authoritative workbook, including all values and formulas."""
import base64, hashlib, json, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
SOURCE=Path(r'C:\Users\320126040\Documents\Codex\AOP2027\AOP 2027 Demand\15926_IT AOP2027 Demand.xlsx')
OUT=ROOT/'output'/'modeler-v3'/'Data file'
OUT.mkdir(parents=True,exist_ok=True)
def safe(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return v.isoformat()
    return v
raw=SOURCE.read_bytes()
values=load_workbook(SOURCE,read_only=True,data_only=True)
formulas=load_workbook(SOURCE,read_only=True,data_only=False)
snapshot={'sourceFile':SOURCE.name,'sourceHash':hashlib.sha256(raw).hexdigest(),'originalBase64':base64.b64encode(raw).decode(),'sheets':[]}
for ws in values:
    fw=formulas[ws.title]
    snapshot['sheets'].append({'name':ws.title,'rows':[[safe(v) for v in row] for row in ws.iter_rows(values_only=True)],'formulas':{cell.coordinate:cell.value for row in fw.iter_rows() for cell in row if cell.data_type=='f'}})
(OUT/'AOP2027-Complete-Workbook-Snapshot.json').write_text(json.dumps(snapshot,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
(OUT/SOURCE.name).write_bytes(raw)
print(json.dumps({'sourceHash':snapshot['sourceHash'],'sheets':[(s['name'],len(s['rows']),len(s['rows'][0])) for s in snapshot['sheets']]}))
