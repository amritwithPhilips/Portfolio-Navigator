from pathlib import Path
import re,sys
from docx import Document
sys.stdout.reconfigure(encoding='utf8');d=Path(__file__).parent
h=(d/'AOP2027-IT-Portfolio-Scenario-Modelling-v5.7.0.html').read_text(encoding='utf8')
for i,m in enumerate(re.finditer(r'<script\b([^>]*)>([\s\S]*?)</script>',h)):
 print('SCRIPT',i,m[1],len(m[2]))
 if i>3:
  (d/f'extra-{i}.js').write_text(m[2],encoding='utf8')
  for x in re.finditer('apply-exco-tradeoff',m[2]):print(m[2][max(0,x.start()-400):x.start()+3200])
root=Path(r'C:\Users\320126040\Documents\Codex\AOP2027\PRD')
for name in ['AOP2027-Modeler-PRD-v5.4.docx','AOP2027-Modeler-Detailed-Requirements-v5.4.docx']:
 doc=Document(root/name);txt='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows]);(d/(name+'.txt')).write_text(txt,encoding='utf8');print('DOC',name,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'chars',len(txt))
print('SUPPORT COMPLETE')
