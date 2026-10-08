from pathlib import Path
import re,json
from docx import Document
from docx.shared import Pt,RGBColor,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Path(__file__).parent;ROOT=Path(r'C:\Users\320126040\Documents\Codex\AOP2027\PRD')
requirements=[]
for line in (D.parent/'finance-plan-20260929/Finance-calculation-proposal.md').read_text(encoding='utf8').splitlines():
 if line.startswith('| FIN-'):
  cells=[x.strip() for x in line.strip('|').split('|')];requirements.append((cells[0],cells[1]))
replacements={
'1.1 Modeler state':None,
'Standalone offline HTML:':'Standalone offline HTML: AOP2027-IT-Portfolio-Scenario-Modelling-v5.7.0.html. Six tabs share one scenario result and one finance allocation ledger. Engine v5.7.0; session schema 4. Source workbooks remain unchanged.',
'Authoritative demand:':'Authoritative Programs source: AOP 2027 Consolidated Live 299262020.xlsx, saved 29 September 2026. The release stores original bytes, cached values, formulas and full SHA-256 provenance.',
'472 demand rows':'433 active Programs; requested demand 285.2129 M€ including 5.19 M€ PMI. Nine missing-request rows and four uninterpretable-decision rows are excluded from active scope and retained in the audit extract.',
'Enterprise-wide charging.':'Enterprise-wide charging at BU row 276 is an allocation key only. Its 40 M€ is excluded from all demand. Apply the user-confirmed 1,000× percentage correction and retain 22.63033175355452% as unallocated. PMI is funded separately on top of the base budget without a BU charge.',
'Operations snapshot':'Operations source: Updated - Preliminary AOP2027 Bill of IT for BU_R_F - 23.Sep.xlsx. All 182 service IDs reconcile to 470.3144630385122 M€; 36 consuming organisations and seven final recipient clusters are retained. No proportional scaling of Operations totals is used.',
}
def update(text):
 for prefix,new in replacements.items():
  if new and text.startswith(prefix):return new
 t=re.sub(r'\b[Pp]rogrammes?\b','Programs',text);t=re.sub(r'\b[Pp]rograms?\b','Programs',t)
 t=t.replace('Version 5.4','Version 5.7.0').replace('Engine v5.2.0','Engine v5.7.0').replace('28 September 2026','30 September 2026').replace('469.7','470.314463')
 t=t.replace('Companion to PRD v5.3','Companion to PRD v5.7.0').replace('AOP 2027 Modeler v5.2.','AOP 2027 Modeler v5.7.0.').replace('AOP2027-Modeler-PRD-v5.2.docx','AOP2027-Modeler-PRD-v5.7.0.docx')
 t=t.replace('AOP2027-IT-Portfolio-Scenario-Modelling-v5.1.html','AOP2027-IT-Portfolio-Scenario-Modelling-v5.7.0.html')
 if t.startswith('Version 5.7.0 supersedes'):t='Version 5.7.0 adds the shared finance allocation ledger, confirmed EW charging split, baseline-capped carry-over rule, source exclusions, scenario request and merit edits, and independent ten-scenario validation. Six-tab navigation and source evidence are retained.'
 if re.search(r'carry.*(ROI|RoI|threshold)|(ROI|RoI).*carry',t,re.I):t='Confirmed carry-over baseline cap; no ROI threshold.'
 t=t.replace('illustrative 100','default 100').replace('illustrative envelope','default planning envelope')
 if re.search(r'hold.*(10.?x|benefit)|(10.?x|benefit).*hold',t,re.I):t='Benefits accepted for prioritization; no benefit-only hold.'
 t=t.replace('10x+ (baseline only)','Confirmed baseline cap').replace('10x+ (baseline)','Confirmed baseline cap').replace('10x (baseline only)','Confirmed baseline cap').replace('10x+','Baseline cap')
 if re.fullmatch(r'\d+(\.\d+)?x',t.strip()):t='Baseline cap'
 t=t.replace('Tighten Carryover Evidence','Balance Sustain and Value').replace('Tighten Carryover','Balance Sustain and Value')
 t=t.replace('472 Programs processed','433 active Programs processed').replace('Prioritization Logic (unchanged from v4.0)','Prioritization Logic (v5.7.0)')
 if t.startswith('Key changes since v4.0: review threshold'):t='Benefit data is accepted for prioritization; no high-ratio review threshold is applied.'
 if t.startswith('BU Own:'):t='Operations own consumption: 320.258096632 M€; Regions: 63.461385361 M€; Functions: 86.594981045 M€.'
 if t.startswith('Total Operations 2027:'):t='Total Operations 2027: 470.314463039 M€; 2026 run-rate: 470.667055448 M€; change: -0.352592410 M€.'
 t=t.replace('Hold; inline review required','Accepted input; no eligibility hold').replace('10.0x (raised from 5.0x in v3.2)','no benefit-ratio threshold')
 if 'scaled to 470.314463' in t:t=t.replace('scaled to 470.314463 M€','reconciled directly to the source 470.3144630385122 M€')
 if t.startswith('Automatic funding holds.'):t='Automatic eligibility retains source decision, condition, name, positive request and merit rules. Benefits alone do not create a hold. Carry Over requires a confirmed approved baseline.'
 t=t.replace('Key changes in v4.0: review threshold at no benefit-ratio threshold.','Benefits are accepted without a high-ratio hold.')
 if 'full reset' in t.lower() and ('Working' in t or 'working' in t or t.startswith('Full reset:')):t='New Working Version creates a draft fork with parent lineage, retaining settings, Programs decisions and evidence reviews.'
 t=t.replace('Applies Optimal Segmentation preset.','New Scenario applies Optimal Segmentation; New Working Version retains the current settings.')
 t=t.replace('Clears all manual overrides (include/exclude decisions).','New Working Version retains manual include/exclude decisions.')
 t=t.replace('Clears all evidence reviews.','New Working Version retains evidence reviews and scenario request/merit edits.')
 t=t.replace('7 BU clusters. Columns: BU Led (M€), Enterprise Wide (M€), Operations (M€), Total.','Seven BU clusters plus Unallocated and PMI. Columns separate Programs types, EW, Operations own/Regions/Functions, PMI and Total.')
 t=t.replace('Organisation and Platform filters with X of Y.','Allocation shows the entire scenario and labels this scope; source-accountability filters apply to the portfolio view.')
 t=t.replace('Allocated AOP ≤ Total AOP x (1 + tolerance%).','Automatic allocation stays within Total AOP and tolerance. Manual additions may exceed it; the over-budget amount remains visible.')
 return t
def paragraphs(doc):
 yield from doc.paragraphs
 for table in doc.tables:
  for row in table.rows:
   for cell in row.cells:yield from cell.paragraphs
 for sec in doc.sections:
  yield from sec.header.paragraphs;yield from sec.footer.paragraphs
for stem in ['AOP2027-Modeler-PRD','AOP2027-Modeler-Detailed-Requirements']:
 doc=Document(ROOT/(stem+'-v5.4.docx'))
 if 'Detailed' in stem:
  deleting=False
  for el in list(doc.element.body):
   text=''.join(el.itertext())
   if 'Finance Requirements (unchanged from v5.0)' in text:deleting=True
   if deleting and 'Acceptance Tests (v5.2 additions)' in text:deleting=False
   if deleting:doc.element.body.remove(el)
  deleting=False
  for el in list(doc.element.body):
   text=''.join(el.itertext())
   if 'UI-10' in text and 'New Working Version' in text and el.tag==qn('w:p'):
    deleting=True
    for p in doc.paragraphs:
     if p._p is el:p.text='UI-10  New Working Version';break
    continue
   if deleting and 'UI-11' in text and el.tag==qn('w:p'):
    deleting=False
    for content in ['Description: Create a draft fork with parent lineage. Retain the current settings, manual Programs decisions, and request/merit/evidence reviews.','Acceptance: Parent scenario is retained; fork is editable; inherited decisions and settings reproduce the parent allocation; revision and event history begin anew. New Scenario is the separate action that resets to Optimal.']:
     from docx.text.paragraph import Paragraph
     node=OxmlElement('w:p');el.addprevious(node);Paragraph(node,doc._body).add_run(content)
   if deleting:doc.element.body.remove(el)
 for p in paragraphs(doc):
  old=p.text;new=update(old)
  new=new.replace('New Working Version reset','New Working Version fork').replace('After reset: Optimal active, no overrides, empty decision log, no trade-off values','Draft fork retains decisions and settings; parent lineage recorded')
  if old!=new:
   if p.runs:p.runs[0].text=new;[setattr(r,'text','') for r in p.runs[1:]]
   else:p.add_run(new)
 for style in doc.styles:
  if style.name in ['Title','Subtitle'] or style.name.startswith('Heading'):
   style.font.color.rgb=RGBColor(0,0,0)
 for p in doc.paragraphs:
  if p.style.name in ['Title','Subtitle'] or p.style.name.startswith('Heading'):
   for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
 doc.add_page_break();doc.add_heading('Finance calculation module — v5.7.0',0)
 doc.add_paragraph('The finance module is the calculation authority for Programs allocation and combined Programs + Operations recipient cost. The source-accountability view and the receiving-BU view represent the same cost and must never be added together.')
 doc.add_heading('Confirmed default and charging key',1)
 doc.add_paragraph('The Optimal default budget is 100 M€. LFR retains the verified share; Carry Over uses approved baseline caps without an ROI threshold; Sustain is 60% of its bucket; Value uses the remaining available budget; IT Simplification is optional and inactive by default. Manual changes retain any overspend visibly. PMI is on top.')
 doc.add_paragraph('EW charge to recipient b = funded EW Programs × charging fraction b. The BU row 276 percentage formulas sum to 0.001 because recipient values totaling 40 are divided by 40,000. The confirmed correction multiplies each source fraction by 1,000. It preserves the exact 22.63033175355452% unallocated share. Display rounding does not change the calculation.')
 for ident,text in requirements:
  if ident=='FIN-05':text='Indirect Cross BU follows its own source allocation, including Not allocated to BU. The requesting Function/Region remains the accountable origin. No fallback to the EW key is applied.'
  if ident=='FIN-08':text='Retain exact source percentage precision. Apply the confirmed 1,000× correction only to an EW vector totaling 0.001; accept a fraction vector totaling 1 directly. Other EW totals are exceptions. BU row percentages are never silently normalized. Missing balances remain Unallocated.'
  if ident=='FIN-15':text='Exclude missing-request records from active Programs overviews, prioritization and exports. Nine such records remain in the refreshed source. Preserve original cells and exclusion reasons. A completed record can return on refresh.'
  if ident=='FIN-19':text='The allocation ledger records source file hash, sheet, row, request/baseline cells, key cell, Programs ID, source and scenario requests, funded amount, recipient fraction and charge. Key cells link to original formulas preserved in the complete source snapshot. The workbook overview identifies scenario and revision. These joins provide an independent manual-style replay.'
  doc.add_heading(ident,2);doc.add_paragraph(update(text))
 doc.add_heading('Reconciliation and known source balances',1)
 doc.add_paragraph('Default funded base Programs = 99.963 M€: EW 49.716 and BU Led 50.247. PMI = 5.190 M€; Operations = 470.3144630385122 M€; total IT = 575.4674630385122 M€. EW unallocated = 11.250895734597165 M€. BU-led unallocated includes 1.357 M€ explicitly assigned to Not allocated and 0.465 M€ missing from Bu-06 percentages. All balances remain included in the total without an invented recipient.')
 doc.add_heading('Scenario acceptance and traceability',1)
 doc.add_paragraph('Ten scenarios cover default Optimal, two Sustain variants, AOP-percent limits, absolute limits, the optional fifth lever, inclusion/partial scope, request/merit edits, F&R/BU ExCo trade-offs and combined 80 M€ stress/recovery. Each scenario must match an independent calculation reading the original two workbooks. Validate all Programs costs, stages, scores and funding, recipient totals, PMI separation, Operations invariance, browser views, persistence and exports.')
 doc.add_paragraph('Trace every allocation through Programs/service ID, source file SHA-256, sheet/row, request cell, baseline cell, key cell, funded amount, recipient fraction and final charge. Technical conformance and source completeness are reported separately. A passing arithmetic result does not resolve the source’s remaining unallocated balances.')
 if stem.endswith('-PRD'):
  t=doc.tables[6]
  replacements_by_row=[['Priority','Name','Classification','Funding order','Default demand M€'],['P1','Legal & Compliance','Compliance name/year evidence, must-do decision, or CCB exception','Verified first; priority tier; merit','53.136'],['P2','Carry Over','Carry Over demand; confirmed funding requires baseline >0 and Must Do/Must Have','Confirmed first; baseline cap','107.5189'],['P3','Sustain & Hygiene','Change & Sustain, Continuous Improvement or Upgrades & Enhancements','Change & Sustain first; then merit','34.367'],['P4','Value Based','Remaining Programs; IT Simplification may be separated as P5','Merit ranking','85.001']]
  for row,vals in zip(t.rows,replacements_by_row):
   for cell,value in zip(row.cells,vals):cell.text=value
  t=doc.tables[8]
  for cell,value in zip(t.rows[2].cells,['2. Carry Over','P2','Confirmed approved baseline','No ROI threshold']):cell.text=value
  t=doc.tables[9];t.rows[0].cells[2].text='P2: Carry Over'
  for row in t.rows[1:]:row.cells[2].text='Confirmed baseline cap'
 for p in doc.paragraphs[:3]:
  for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
 for style in doc.styles:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
 for table in doc.tables:
  for ri,row in enumerate(table.rows):
   trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
   if ri==0:trPr.append(OxmlElement('w:tblHeader'))
   for cell in row.cells:
    for para in cell.paragraphs:
     for run in para.runs:run.font.name='Calibri';run.font.size=Pt(9);run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
  borders=OxmlElement('w:tblBorders')
  for side in ['top','left','bottom','right','insideH','insideV']:
   edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
  table._tbl.tblPr.append(borders)
 doc.save(D/(stem+'-v5.7.0.docx'))
print('Two updated documents saved')
