"""Independent workbook-to-scenario replay; does not import the application engine."""
import json,re,math,hashlib,collections,sys
from pathlib import Path
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding='utf-8');D=Path(__file__).parent
ROOT=Path(r'C:\Users\320126040\Documents\Codex\AOP2027\AOP 2027 Demand')
book=ROOT/'AOP 2027 Consolidated Live 299262020.xlsx';wb=load_workbook(book,data_only=True,read_only=True)
scenarios=json.loads((D/'audit/scenarios.json').read_text(encoding='utf8'))['scenarios'];checks=[];mismatches=[]
def clean(v):return re.sub(r'\s+',' ',str(v).replace('\u200b',' ').replace('\xa0',' ')).strip() if v is not None else ''
def key(v):return clean(v).lower().replace('—','-').replace('–','-')
def num(v):
 if isinstance(v,(float,int)) and not isinstance(v,bool):return v
 if isinstance(v,str) and re.fullmatch(r'-?\d+(\.\d+)?',v.strip()):return float(v)
 return None
def rnd(v):return math.floor(v*1e6+.50000001)/1e6
def check(n,ok,detail=''):checks.append({'name':n,'passed':bool(ok),'detail':detail})
records=[];excluded=[];charging=None
for sn in ['EW','BU']:
 rows=list(wb[sn].iter_rows(values_only=True));hi=next(i for i,r in enumerate(rows[:10]) if any(key(v)=='item no.' for v in r));headers=[key(v) for v in rows[hi]]
 def ix(label):return headers.index(label)
 for rownum,row in enumerate(rows[hi+1:],hi+2):
  def val(label):return row[ix(label)]
  item=clean(val('item no.'));name=clean(val('initiative name'));request=num(val('it investment requested 2027 [k€]'))
  if (not item or re.fullmatch('0+',item)) and not name and request is None:continue
  r={'id':item if item and not re.fullmatch('0+',item) else f'{sn}-row-{rownum}','name':name,'stream':'ew' if sn=='EW' else 'bu','sourceRow':rownum,'sourceSheet':sn,'type':clean(val('type')),'cost':None if request is None else rnd(request/1000),'newCarry':clean(val('new initiative or carry-over')),'category':clean(val('category')),'priority':clean(val('priority')),'overall':clean(val('overall aop decision')),'band':clean(val('waste band')),'lfrName':clean(val('lfr & compliancy (name)')),'lfrYear':clean(val('lfr & compliancy year')),'baselined':num(val('carry over baselined')),'benefit':num(val('quantified benefits (in k€)')),'benefitDetails':clean(val('qualitative benefit')),'impact':clean(val('impact and cost of not doing')),'problem':clean(val('problem statement')),'scope':clean(val('scope')),'payback':num(val('payback period')),'remarks':clean(row[headers.index('remarks / evidence')]) if 'remarks / evidence' in headers else ''}
  r['baselined']=rnd(r['baselined']/1000) if r['baselined'] is not None else None;r['benefit']=rnd(r['benefit']/1000) if r['benefit'] is not None else None
  r['alloc']={clean(rows[hi][i]).removesuffix(' %').strip():row[i] for i,h in enumerate(headers) if h.endswith('%') and 'capitalization' not in h and isinstance(row[i],(float,int))}
  if key(name)=='enterprise-wide charging':charging=r;continue
  if r['cost'] is None or r['overall'].startswith('#') or r['band'].startswith('#') or r['overall']=='0':excluded.append(r);continue
  records.append(r)
check('Independent active source count',len(records)==433,str(len(records)));check('Source key total and scale',abs(sum(charging['alloc'].values())-.001)<1e-12)
harm=r'\b(cannot|unable|not able|will (?:not|lose|fail)|loss of|revenue (?:loss|at risk)|penalt\w*|non[- ]?complian\w*|end of (?:life|support)|out of support|downtime|disrupt\w*|delayed|forfeit|no longer|security breach\w*)\b'
affected=r'\b(customer\w*|patient\w*|user\w*|business|product\w*|system\w*|platform\w*|operation\w*|revenue|audit\w*|compliance|manufactur\w*|service\w*|support|process\w*)\b'
specific=r'\b(cost (avoidance|reduction|saving)|efficien(cy|t)|automat(e|ion|ing)|eliminat(e|ion|ing)|compliance|regulatory|standardiz|consolidat|streamlin|simplif(y|ication)|reduc(e|tion|ing) (time|effort|risk|cost|manual)|enabl(e|ement|ing)|avoid(ance|ing)|retir(e|ement|ing)|decommission|migrat(e|ion|ing)|modern|optimi[sz]|improv(e|ement|ing) (process|quality|speed|accuracy))\b'
def match(pattern,text):return bool(re.search(pattern,text,re.I))
def enrich(r,settings,review):
 r=dict(r);r['cost']=review.get('requested',r['cost']);cost=r['cost'];benefit=r['benefit'];ratio=benefit/cost if cost>0 and benefit is not None else None
 real=lambda x:bool(x) and not re.fullmatch(r'0|na|n/a|tbd|none|not applicable',key(x))
 st=0 if r['id'] in ['EW-11','EW-65','EW-71'] or real(r['lfrName']) or real(r['lfrYear']) or 'must-do' in key(r['overall']) else 1 if key(r['newCarry'])=='carry over' else 2 if key(r['newCarry'])=='change & sustain' or key(r['category']) in ['continuous improvement','continous improvement','continuous improvement (ci)','upgrades & enhancements'] else 3
 if st==3 and settings['itsimpMode'] and match('decomm|replatform',r['category']):st=4
 text=r['benefitDetails'];bs='quantitative'
 if ratio is not None:bp=0 if ratio<1 else 15 if ratio<3 else 25 if ratio<5 else 35 if ratio<7 else 50
 elif match(r'\b(revenue (increase|growth|generation)|productivity (saving|gain|improvement)|sales growth|top[- ]?line growth)\b',text):bp=0;bs='unquantified'
 elif len(text)<10 or re.fullmatch(r'na|tbd|n/a|none|soft benefit|not applicable|to be determined',text,re.I):bp=0;bs='none'
 elif match(r'\d[\d,.]*\s*(%|fte|hours?|days?|processes|sites?|users?|systems?|applications?|minutes?|teams?)\b',text) or (match(r'\b(enterprise|cross[- ]?functional|global|company[- ]?wide|all (bus|regions|markets|functions))\b',text) and match(specific,text)):bp=30;bs='qualitative'
 elif match(specific,text) and len(text)>50:bp=20;bs='qualitative'
 else:bp=10;bs='qualitative'
 evidence=' '.join(x for x in [r['impact'],r['problem'],r['scope']] if x);cp=5 if evidence else 0
 if match(harm,evidence) and match(affected,evidence):cp=15 if match(r'\b(global|enterprise|company[- ]wide|worldwide|multi[- ]region|all regions|all markets)\b',evidence) or match(r'\d[\d,.]*\s*(%|fte|hours?|days?|k€|m€|eur|million|sites?)\b',evidence) or match(r'(?:€|EUR)\s*\d',evidence) else 10
 if review.get('consequence') in [0,5,10,15] and evidence and clean(review.get('note')):cp=review['consequence']
 pp=25 if key(r['priority'])=='must do' else 18 if key(r['priority'])=='must have' else 5 if key(r['priority']).startswith('deprior') else 0
 dp=10 if key(r['band'])=='healthy' else 6 if key(r['band'])=='contain' else 0;pb=r['payback'];pk=0 if pb is None or pb<=0 else 10 if pb<=1 else 8 if pb<=2 else 6 if pb<=3 else 4 if pb<=5 else 0
 components=[v/m*w for v,m,w in zip([bp,pp,cp,dp,pk],[50,25,15,10,10],settings['weights'])];score=review.get('merit',rnd(sum(components)))
 confirmed=st==1 and (r['baselined'] or 0)>0 and key(r['priority']) in ['must do','must have'];cap=min(cost,r['baselined']) if st==1 and (r['baselined'] or 0)>0 else cost
 decision=key(r['overall']);held=not r['name'] or cost<=0 or 'priority tbd' in decision
 if decision and not re.fullmatch(r'healthy\s*-\s*rank in priority matrix|healthy\s*/\s*must-do\s*-\s*(p1\s*-\s*immediate|p2\s*-\s*scheduled|priority tbd)|contain\s*-\s*conditions',decision):held=True
 if (key(r['band'])=='contain' or decision.startswith('contain')) and not clean(r['remarks']) and not clean(review.get('condition')) and st!=0:held=True
 if st==1 and not confirmed:held=True
 if st==3 and score<settings['merit']:held=True
 r.update(stage=st,score=score,components=components,ratio=ratio,confirmed=confirmed,cap=cap,held=held,verified=(real(r['lfrName']) and real(r['lfrYear'])) or r['id'] in ['EW-11','EW-65','EW-71'])
 return r
def order(r):
 st=r['stage'];score=-r['score']
 if st==0:return (st,0 if r['verified'] else 1,0 if 'p1' in key(r['overall']) else 1 if 'p2' in key(r['overall']) else 2 if r['id'] in ['EW-11','EW-65','EW-71'] else 3,score)
 if st==1:return (st,0 if r['confirmed'] else 1 if r['ratio'] is not None else 2,-(r['ratio'] if r['ratio'] is not None else -float('inf')),score)
 if st==2:return (st,0 if key(r['newCarry'])=='change & sustain' else 1,score)
 return (st,score,tuple(int(v) if v.isdigit() else v.lower() for v in re.split('(\d+)',r['id'])))
mp={'Imaging Informatics':'EI','Clinical Informatics':'EI','EI Other':'EI','AMD':'Connected Care','HPM':'Connected Care','S&RC':'Connected Care','I&D':'Innovation & Design','IP and S':'Int. Property & Standards','IGT-D':'Image Guided Therapy','IGT-S':'Image Guided Therapy','MR':'Precision Diagnosis','CT':'Precision Diagnosis','DXR':'Precision Diagnosis','PD Other':'Precision Diagnosis','US':'Precision Diagnosis','G and B':'Personal Health','MCC':'Personal Health','OHC':'Personal Health','PH Other':'Personal Health','Not allocated':'Unallocated'}
opfile=ROOT/'Updated - Preliminary AOP2027 Bill of IT for BU_R_F - 23.Sep.xlsx';op=load_workbook(opfile,data_only=True,read_only=True)
tco=list(op['Comparable AOP2027'].iter_rows(values_only=True));service=list(op['AOP Svc. Costs (21.09)'].iter_rows(values_only=True));bill=list(op['AOP2027 Bill of IT'].iter_rows(values_only=True))
service_total=sum(r[6]/1000 for r in service[1:] if r[1] is not None and num(r[6]) is not None)
consumption=sum(v/1000 for r in bill[5:180] if r[1] is not None and num(r[12]) is not None for v in r[123:159] if isinstance(v,(int,float)))
check('Operations all service costs',abs(service_total-470.3144630385122)<1e-8);check('Operations 36 recipient consumption columns',abs(consumption-service_total)<1e-8,str(consumption))
for case in scenarios:
 s=case['settings'];rs=[enrich(r,s,case['reviews'].get(r['id'],{})) for r in records];base=sorted([r for r in rs if key(r['type'])!='pmi'],key=order)
 dem=[rnd(sum(r['cost'] for r in base if r['stage']==st)) for st in range(5)]
 def target(k,st):
  q=s[k];return min(dem[st],q.get('amount',0) if q['basis']=='amount' else (dem[st] if q['basis']=='bucket' else s['budget'])*q['pct']/100)
 targets=[target('lfr',0),float('inf'),target('sustain',2),target('value',3),target('itsimp',4) if s['itsimpMode'] else 0];spent=[0]*5;auto=0
 for r in base:
  st=r['stage'];residual=max(0,rnd(s['budget']*(1+s['tolerance']/100)-sum(spent[:3])))
  limit=min(max(0,residual-min(targets[4],dem[4])),targets[3]*(1+s['tolerance']/100)) if st==3 else min(dem[st],targets[st]*(1+s['tolerance']/100)) if st!=1 else float('inf')
  payable=r['cap'] if st==1 else r['cost'];fund=0
  if not r['held'] and spent[st]+payable<=limit+1e-9 and auto+payable<=s['budget']*(1+s['tolerance']/100)+1e-9:fund=payable;spent[st]+=fund;auto+=fund
  o=case['overrides'].get(r['id']);r['automatic']=fund
  if o and o.get('state')!='needs-review':fund=0 if o['choice']=='exclude' else rnd(o.get('amount',r['cost']))
  r['final']=fund
 for r in rs:
  if key(r['type'])=='pmi':r['automatic']=r['final']=r['cost']
 actual={r['id']:r for r in case['result']['rows']};errors=[]
 for r in rs:
  x=actual.get(r['id'])
  if not x:errors.append((r['id'],'missing'));continue
  for k in ['cost','stage','score','automatic','final']:
   if abs(r[k]-x[k])>1e-6:errors.append((r['id'],k,r[k],x[k]))
 check(case['id']+' independent record funding and merit',not errors,errors[:12]);mismatches.extend(errors)
 expected=collections.defaultdict(float)
 for r in rs:
  if key(r['type'])=='pmi':expected['PMI on top']+=r['final'];continue
  isew=r['stream']=='ew' or match('type iii|enterprise wide',r['type']);alloc=charging['alloc'] if isew else r['alloc'];scale=1000 if isew else 1;fractions={k:v*scale for k,v in alloc.items()}
  if sum(fractions.values())>1+1e-8:expected['Unallocated']+=r['final'];continue
  for n,v in fractions.items():expected[mp.get(n,'Unallocated')]+=r['final']*v
  remainder=r['final']*(1-sum(fractions.values()))
  if remainder>1e-9:expected['Unallocated']+=remainder
 for row in tco[3:30]:
  if not str(row[1]).startswith('-'):expected[str(row[1]).strip()]+=row[6]+row[7]+row[8]
 errors=[(r['recipient'],expected[r['recipient']],r['total']) for r in case['result']['finance']['rows'] if abs(expected[r['recipient']]-r['total'])>1e-8]
 check(case['id']+' independent receiving organisation totals',not errors,errors)
 check(case['id']+' exact source percentage lineage',all(e['pillar']!='Programs' or not e.get('keyCell') or abs(next(r for r in rs if r['id']==e['id'])['final']*e['fraction']-e['amount'])<1e-8 for e in case['result']['finance']['ledger']))
check('Source workbook hash stable',hashlib.sha256(book.read_bytes()).hexdigest()=='70480f321fa60e6cba983bf63e72572570b2821ec0a222b604449aafe9f4f470')
(D/'audit/oracle-qa.json').write_text(json.dumps({'checks':checks,'recordsPerScenario':len(records),'scenarios':len(scenarios),'fieldComparisons':len(records)*len(scenarios)*5,'mismatches':mismatches},indent=2),encoding='utf8')
print(json.dumps({'checks':len(checks),'passed':sum(c['passed'] for c in checks),'failures':[c for c in checks if not c['passed']]},indent=2))
