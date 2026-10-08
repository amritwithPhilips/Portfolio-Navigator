from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
D=Path(__file__).parent;req=json.loads((D/'requirements.json').read_text(encoding='utf8'))
def new(title):
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8)
 for n in ['Normal','Title','Heading 1','Heading 2','Header','Footer']:
  st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 for st in d.styles:
  for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Normal'].paragraph_format.line_spacing=1.12;d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 1'].paragraph_format.space_before=Pt(12);d.styles['Heading 1'].paragraph_format.space_after=Pt(6)
 s.header.paragraphs[0].text='AOP 2027  |  ExCo decision workspace  |  Version 5.9.0';f=s.footer.paragraphs[0];f.text='30 September 2026  |  ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field);d.add_paragraph(title,'Title');return d
def p(d,t):d.add_paragraph(t)
def h(d,t):d.add_paragraph(t,'Heading 1')
d=new('AOP 2027 ExCo workspace product requirements')
p(d,'Version 5.9.0 adds leader-specific decisions, funding opportunities and customizable evidence to the ExCo view. This addendum supplements the v5.8.0 ExCo specification and the unchanged v5.7.0 detailed modeler specification. Both earlier HTML builds and source workbooks remain unchanged.')
h(d,'Leader decisions')
p(d,'The Decisions workspace supports all base Programs priorities, with PMI retained separately on top. Leaders can include or exclude selected Programs, return them to automatic funding, and review requested investment, merit and partial scope. Every funding change is previewed before apply. Reviewed merit is optional; a blank value uses the calculated score. Rationale is required for individual reviews.')
h(d,'Charts and evidence')
p(d,'The decision map toggles between requested Demand and quantified Hard Benefits against merit. Bubble size reflects requested investment; colour distinguishes funded and not funded. Clicking or pressing Enter on a bubble opens source evidence. Hard Benefits excludes soft benefit claims.')
p(d,'Custom columns include normalized scenario fields and original EW and BU source columns. Leaders can search, add or remove fields while Programs identity and source access remain visible. Layouts are stored separately by leader in the saved scenario, retained through reload, fork and portable export. A fresh scenario uses default layouts.')
h(d,'Shared company state')
p(d,'ExCo leader selection synchronizes the Programs filters in the detailed workbench. Request, merit, inclusion and funding edits work in both directions. All six views and exported finance ledgers use the same current result. Receiving-organisation cost allocation remains company-wide, distinct from accountable Programs ownership.')
p(d,'Company funding settings expose the base budget, tolerance, merit floor, all five score weights, editable priority lever bases and optional IT Simplification. These controls affect all leaders and are explicitly labelled company-wide. Preview shows the resulting company and recipient costs.')
d.add_page_break();h(d,'Additional funding calculation')
p(d,'The user enters a positive additional company envelope, one ExCo leader, a stream or both streams, and a rationale. The planner uses current scenario requests, merit and source eligibility. It ranks remaining eligible scopes by priority stage, then descending merit, ascending request and ID. It selects whole remaining scopes that fit the additional envelope.')
p(d,'Confirmed carry-over remains capped at the approved baseline. Explicit manual exclusions and held Programs are not selected automatically. The added envelope increases the company budget by the entered amount. Any unused extra remains headroom. Existing allocations across all leaders are retained, so increasing one leader’s funding does not silently refill other portfolios.')
p(d,'Retained automatic allocations become explicit recorded overrides. A separate preview can return those retained decisions to automatic, while Undo restores the entire previous top-up transaction. The existing finalisation and changed-source decision holds continue to apply.')
h(d,'Benefit interpretation')
p(d,'Preview reports newly funded productivity claims as a savings proxy, separate revenue claims and separate soft claims. All are submitted lifetime claims. The model does not infer annual cash savings or incremental benefits from additional funding of an already funded partial scope. Missing quantified evidence is not converted into a realised return.')
h(d,'Validation and finance continuity')
p(d,'The release validates 32 existing scenarios, 144 leader and stream combinations at €1M, €5M and €15M extra budgets, source-linked Excel exports and independent workbook replay. The final QA report contains assertion counts, source reconciliation, UI checks and the tested additional-funding example.')
p(d,'Default Optimal remains €100M, with €99.963M funded base Programs and €575.467463M total IT. The approved 22.6303317536% EW unallocated share is retained. BU row 276 is a percentage charging key; its €40M remains excluded from demand. PMI and Operations treatment remain unchanged.')
d.save(D/'AOP2027-Modeler-PRD-v5.9.0-ExCo-Addendum.docx')
d=new('AOP 2027 ExCo workspace detailed requirements')
p(d,'Version 5.9.0 acceptance contract. These requirements supplement the v5.8.0 and v5.7.0 specifications. Validation evidence is in ExCo-QA-report-v5.9.0.html and the accompanying audit files.')
for ident,title,text in req:
 if ident=='EXCO59-06':d.add_page_break()
 h(d,ident+'  '+title);p(d,text)
d.save(D/'AOP2027-Modeler-Detailed-Requirements-v5.9.0-ExCo-Addendum.docx');print('Created both specification addenda')
